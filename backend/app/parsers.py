from __future__ import annotations

import csv
import hashlib
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import Anchor, AnchorKind, DatasetRole, Record, Source
from .semantics import canonicalize_row, infer_role, infer_semantic_map, infer_text_role


def stable_id(prefix: str, *parts: Any) -> str:
    payload = "|".join(str(part) for part in parts)
    return f"{prefix}_{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:20]}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    return value


def detect_encoding(path: Path) -> str:
    # GDPdU exports often remain ASCII for thousands of rows before the first
    # German umlaut. Sampling only the file head can therefore misclassify a
    # CP1252 export as UTF-8 and fail late, silently removing an entire ledger.
    sample = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            sample.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "latin-1"


def detect_media_type(path: Path) -> str:
    extension_map = {
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pdf": "application/pdf",
        ".csv": "text/csv",
        ".xml": "application/xml",
        ".dtd": "application/xml-dtd",
    }
    return extension_map.get(path.suffix.lower(), mimetypes.guess_type(path.name)[0] or "application/octet-stream")


@dataclass
class ParsedDataset:
    dataset_key: str
    role: DatasetRole
    headers: list[str]
    semantic_map: dict[str, str]
    records: list[Record] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParseBundle:
    source: Source
    datasets: list[ParsedDataset] = field(default_factory=list)
    anchors: list[Anchor] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class DossierParser:
    def __init__(self, case_id: str, originals_root: Path) -> None:
        self.case_id = case_id
        self.root = originals_root
        self.gdpdu_schemas = self._read_gdpdu_schemas()

    def _read_gdpdu_schemas(self) -> dict[Path, dict[str, list[str]]]:
        schemas: dict[Path, dict[str, list[str]]] = {}
        for index_path in self.root.rglob("index.xml"):
            try:
                root = ET.parse(index_path).getroot()
            except (ET.ParseError, OSError):
                continue
            directory_schema: dict[str, list[str]] = {}
            for table in root.findall(".//Table"):
                url = table.findtext("URL")
                columns = [node.text or "" for node in table.findall(".//VariableColumn/Name")]
                if url and columns:
                    directory_schema[url] = columns
            if directory_schema:
                schemas[index_path.parent.resolve()] = directory_schema
        return schemas

    def parse_all(self) -> list[ParseBundle]:
        bundles: list[ParseBundle] = []
        for path in sorted(p for p in self.root.rglob("*") if p.is_file() and p.name != ".DS_Store"):
            bundles.append(self.parse(path))
        return bundles

    def parse(self, path: Path) -> ParseBundle:
        relative = path.relative_to(self.root).as_posix()
        digest = sha256_file(path)
        source_id = stable_id("src", self.case_id, relative, digest)
        source = Source(
            source_id=source_id,
            case_id=self.case_id,
            name=path.name,
            relative_path=relative,
            sha256=digest,
            media_type=detect_media_type(path),
            size=path.stat().st_size,
            provenance_family=stable_id("pf", self.case_id, relative),
        )
        bundle = ParseBundle(source=source)
        try:
            suffix = path.suffix.lower()
            if suffix in {".csv", ".txt"}:
                self._parse_delimited(path, bundle)
            elif suffix == ".xlsx":
                self._parse_xlsx(path, bundle)
            elif suffix == ".pdf":
                self._parse_pdf(path, bundle)
            elif suffix == ".docx":
                self._parse_docx(path, bundle)
            elif suffix == ".xml":
                self._parse_xml(path, bundle)
            elif suffix == ".dtd":
                self._parse_text(path, bundle, AnchorKind.TEXT)
            else:
                bundle.source.status = "unsupported"
                bundle.warnings.append(f"Unsupported file type: {suffix or 'unknown'}")
            if suffix == ".xlsx" and not any(dataset.records for dataset in bundle.datasets):
                bundle.warnings.append("Empty workbook: no populated records; treated as missing evidence, not fraud.")
        except Exception as exc:  # keep ingestion complete and make failures visible
            bundle.source.status = "error"
            bundle.warnings.append(f"{type(exc).__name__}: {exc}")

        roles = sorted({dataset.role for dataset in bundle.datasets} | ({infer_text_role(" ".join(a.quote for a in bundle.anchors))} if bundle.anchors else set()), key=str)
        bundle.source.roles = [role for role in roles if role != DatasetRole.UNKNOWN] or [DatasetRole.UNKNOWN]
        bundle.source.record_count = sum(len(dataset.records) for dataset in bundle.datasets)
        bundle.source.metadata = {
            "datasets": [
                {
                    "dataset_key": dataset.dataset_key,
                    "role": dataset.role,
                    "headers": dataset.headers,
                    "semantic_map": dataset.semantic_map,
                    **dataset.metadata,
                }
                for dataset in bundle.datasets
            ],
            "warnings": bundle.warnings,
        }
        return bundle

    def _anchor(
        self,
        bundle: ParseBundle,
        kind: AnchorKind,
        locator: dict[str, Any],
        quote: str,
        raw_value: Any = None,
        normalized_value: Any = None,
        row_hash: str | None = None,
    ) -> Anchor:
        anchor_id = stable_id("anc", bundle.source.source_id, kind, locator)
        anchor = Anchor(
            anchor_id=anchor_id,
            case_id=self.case_id,
            source_id=bundle.source.source_id,
            kind=kind,
            locator=locator,
            quote=quote[:8_000],
            raw_value=raw_value,
            normalized_value=normalized_value,
            row_hash=row_hash,
        )
        bundle.anchors.append(anchor)
        return anchor

    def _record(
        self,
        bundle: ParseBundle,
        dataset: ParsedDataset,
        row_number: int,
        row: dict[str, Any],
        anchor: Anchor,
    ) -> Record:
        serializable = {str(key): json_value(value) for key, value in row.items()}
        raw_hash = hashlib.sha256(repr(serializable).encode("utf-8")).hexdigest()
        return Record(
            record_id=stable_id("rec", bundle.source.source_id, dataset.dataset_key, row_number, raw_hash),
            case_id=self.case_id,
            source_id=bundle.source.source_id,
            dataset_key=dataset.dataset_key,
            role=dataset.role,
            row_number=row_number,
            data=serializable,
            canonical=canonicalize_row(serializable, dataset.role, dataset.semantic_map),
            anchor_id=anchor.anchor_id,
            raw_hash=raw_hash,
        )

    def _parse_delimited(self, path: Path, bundle: ParseBundle) -> None:
        encoding = detect_encoding(path)
        schema = self.gdpdu_schemas.get(path.parent.resolve(), {}).get(path.name)
        with path.open("r", encoding=encoding, newline="") as handle:
            sample = handle.read(16_384)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
                delimiter, quotechar = dialect.delimiter, dialect.quotechar or '"'
            except csv.Error:
                delimiter, quotechar = ";", '"'
            rows = list(csv.reader(handle, delimiter=delimiter, quotechar=quotechar))
        if not rows:
            return

        if schema:
            headers = schema
            data_rows = rows
            first_data_line = 1
            header_mode = "gdpdu-index"
        else:
            likely_header = self._looks_like_header(rows[0])
            headers = [cell.strip() or f"column_{idx + 1}" for idx, cell in enumerate(rows[0])]
            if not likely_header:
                headers = [f"column_{idx + 1}" for idx in range(len(rows[0]))]
            data_rows = rows[1:] if likely_header else rows
            first_data_line = 2 if likely_header else 1
            header_mode = "native" if likely_header else "generated"

        semantic_map = infer_semantic_map(headers)
        role = infer_role(headers, path.stem, semantic_map)
        dataset = ParsedDataset(
            dataset_key=path.name,
            role=role,
            headers=headers,
            semantic_map=semantic_map,
            metadata={"encoding": encoding, "delimiter": delimiter, "header_mode": header_mode},
        )
        for offset, cells in enumerate(data_rows):
            if not any(str(cell).strip() for cell in cells):
                continue
            line_number = first_data_line + offset
            padded = cells + [""] * max(0, len(headers) - len(cells))
            row = dict(zip(headers, padded[: len(headers)], strict=True))
            raw_hash = hashlib.sha256(delimiter.join(cells).encode(encoding, errors="replace")).hexdigest()
            quote = " · ".join(f"{key}={value}" for key, value in row.items() if str(value).strip())
            anchor = self._anchor(
                bundle,
                AnchorKind.DELIMITED,
                {"record": line_number, "fields": headers},
                quote,
                raw_value=row,
                row_hash=raw_hash,
            )
            record = self._record(bundle, dataset, line_number, row, anchor)
            anchor.normalized_value = record.canonical
            dataset.records.append(record)
        bundle.datasets.append(dataset)

    @staticmethod
    def _looks_like_header(row: list[str]) -> bool:
        nonempty = [cell.strip() for cell in row if cell.strip()]
        if len(nonempty) < 2 or len(set(nonempty)) != len(nonempty):
            return False
        identifier_like = sum(bool(re.search(r"[A-Za-zÄÖÜäöü_]", cell)) for cell in nonempty)
        numeric_like = sum(bool(re.fullmatch(r"[-+]?\d+[.,]?\d*", cell)) for cell in nonempty)
        return identifier_like >= max(2, len(nonempty) // 2) and numeric_like == 0

    def _parse_xlsx(self, path: Path, bundle: ParseBundle) -> None:
        workbook = load_workbook(path, read_only=True, data_only=False)
        displayed_workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            values = [list(row) for row in sheet.iter_rows(values_only=True)]
            displayed_sheet = displayed_workbook[sheet.title]
            displayed_values = [list(row) for row in displayed_sheet.iter_rows(values_only=True)]
            if not values:
                continue
            header_index = self._find_header_row(values[:12])
            headers = [str(value).strip() if value is not None and str(value).strip() else f"column_{idx + 1}" for idx, value in enumerate(values[header_index])]
            title = " ".join(str(value) for row in values[: header_index + 1] for value in row if value is not None)
            semantic_map = infer_semantic_map(headers)
            role = infer_role(headers, f"{sheet.title} {title}", semantic_map)
            dataset = ParsedDataset(
                dataset_key=sheet.title,
                role=role,
                headers=headers,
                semantic_map=semantic_map,
                metadata={"sheet": sheet.title, "header_row": header_index + 1},
            )
            for zero_index, cells in enumerate(values[header_index + 1 :], start=header_index + 1):
                if not any(value is not None and str(value).strip() for value in cells):
                    continue
                padded = cells + [None] * max(0, len(headers) - len(cells))
                shown = displayed_values[zero_index] if zero_index < len(displayed_values) else []
                shown = shown + [None] * max(0, len(headers) - len(shown))
                effective = [
                    shown[idx]
                    if isinstance(padded[idx], str) and padded[idx].startswith("=") and shown[idx] is not None
                    else padded[idx]
                    for idx in range(len(headers))
                ]
                row = {headers[idx]: json_value(effective[idx]) for idx in range(len(headers))}
                formulas = {
                    f"{get_column_letter(idx + 1)}{zero_index + 1}": padded[idx]
                    for idx in range(len(headers))
                    if isinstance(padded[idx], str) and padded[idx].startswith("=")
                }
                displayed_row = {headers[idx]: json_value(shown[idx]) for idx in range(len(headers))}
                excel_row = zero_index + 1
                end_col = get_column_letter(len(headers))
                quote = " · ".join(f"{key}={value}" for key, value in row.items() if value not in (None, ""))
                raw_hash = hashlib.sha256(repr(row).encode("utf-8")).hexdigest()
                anchor = self._anchor(
                    bundle,
                    AnchorKind.XLSX,
                    {"sheet": sheet.title, "range": f"A{excel_row}:{end_col}{excel_row}"},
                    quote,
                    raw_value={"values": row, "formulas": formulas, "displayed": displayed_row},
                    row_hash=raw_hash,
                )
                record = self._record(bundle, dataset, excel_row, row, anchor)
                anchor.normalized_value = record.canonical
                dataset.records.append(record)
            bundle.datasets.append(dataset)

    @staticmethod
    def _find_header_row(rows: list[list[Any]]) -> int:
        best_index, best_score = 0, -1
        for index, row in enumerate(rows):
            nonempty = [value for value in row if value is not None and str(value).strip()]
            string_count = sum(isinstance(value, str) for value in nonempty)
            unique_count = len({str(value) for value in nonempty})
            score = len(nonempty) * 10 + string_count + unique_count - index
            if len(nonempty) >= 2 and score > best_score:
                best_index, best_score = index, score
        return best_index

    def _parse_pdf(self, path: Path, bundle: ParseBundle) -> None:
        all_text: list[str] = []
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                words = page.extract_words(use_text_flow=True, keep_blank_chars=False)
                lines: list[list[dict[str, Any]]] = []
                for word in words:
                    if not lines or abs(float(word["top"]) - float(lines[-1][0]["top"])) > 3:
                        lines.append([word])
                    else:
                        lines[-1].append(word)
                for line_number, line in enumerate(lines, start=1):
                    line.sort(key=lambda item: float(item["x0"]))
                    text = " ".join(str(item["text"]) for item in line).strip()
                    if not text:
                        continue
                    all_text.append(text)
                    bbox = [
                        min(float(item["x0"]) for item in line),
                        min(float(item["top"]) for item in line),
                        max(float(item["x1"]) for item in line),
                        max(float(item["bottom"]) for item in line),
                    ]
                    self._anchor(bundle, AnchorKind.PDF, {"page": page_number, "line": line_number, "bbox": bbox}, text)

        if not all_text:
            self._ocr_pdf(path, bundle)
            all_text = [anchor.quote for anchor in bundle.anchors]
        role = infer_text_role("\n".join(all_text))
        dataset = ParsedDataset(
            path.name,
            role,
            ["label", "amount"] if role == DatasetRole.FINANCIAL_STATEMENT else [],
            {"text": "label", "amount": "amount"} if role == DatasetRole.FINANCIAL_STATEMENT else {},
            metadata={"pages": max((a.locator.get("page", 0) for a in bundle.anchors), default=0)},
        )
        if role == DatasetRole.FINANCIAL_STATEMENT:
            for row_number, anchor in enumerate(bundle.anchors, start=1):
                match = re.match(r"^(.*?)\s+(-?[0-9][0-9.]*,[0-9]{2})$", anchor.quote)
                if not match:
                    continue
                row = {"label": match.group(1).strip(), "amount": match.group(2)}
                record = self._record(bundle, dataset, row_number, row, anchor)
                anchor.normalized_value = record.canonical
                dataset.records.append(record)
        bundle.datasets.append(dataset)

    def _ocr_pdf(self, path: Path, bundle: ParseBundle) -> None:
        with tempfile.TemporaryDirectory(prefix="ledgerlens-ocr-") as temp_dir:
            prefix = Path(temp_dir) / "page"
            subprocess.run(["pdftoppm", "-png", "-r", "200", str(path), str(prefix)], check=True, capture_output=True)
            for page_number, image in enumerate(sorted(Path(temp_dir).glob("page-*.png")), start=1):
                result = subprocess.run(
                    ["tesseract", str(image), "stdout", "-l", "deu+eng", "tsv"],
                    check=True,
                    capture_output=True,
                    text=True,
                )
                for line_number, line in enumerate(result.stdout.splitlines()[1:], start=1):
                    cells = line.split("\t")
                    if len(cells) < 12 or not cells[11].strip():
                        continue
                    left, top, width, height = map(int, cells[6:10])
                    self._anchor(
                        bundle,
                        AnchorKind.PDF,
                        {"page": page_number, "line": line_number, "bbox_px": [left, top, left + width, top + height], "ocr": True},
                        cells[11].strip(),
                    )

    def _parse_docx(self, path: Path, bundle: ParseBundle) -> None:
        document = Document(path)
        rendered_locations = self._render_docx_locations(path)

        def locator_for(text: str, base: dict[str, Any]) -> dict[str, Any]:
            normalized = re.sub(r"\s+", " ", text).strip().lower()
            match = next(
                (
                    item
                    for item in rendered_locations
                    if normalized in item["text"].lower() or item["text"].lower() in normalized
                ),
                None,
            )
            return {
                **base,
                "character_range": [0, len(text)],
                "rendered_page": match["page"] if match else None,
                "rendered_bbox": match["bbox"] if match else None,
            }

        text_parts: list[str] = []
        for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            text_parts.append(text)
            self._anchor(bundle, AnchorKind.DOCX, locator_for(text, {"paragraph": paragraph_index}), text)
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                for column_index, cell in enumerate(row.cells, start=1):
                    text = cell.text.strip()
                    if text:
                        text_parts.append(text)
                        self._anchor(
                            bundle,
                            AnchorKind.DOCX,
                            locator_for(text, {"table": table_index, "row": row_index, "column": column_index}),
                            text,
                        )
        role = infer_text_role("\n".join(text_parts))
        bundle.datasets.append(ParsedDataset(path.name, role, [], {}, metadata={"paragraphs": len(document.paragraphs)}))

    @staticmethod
    def _render_docx_locations(path: Path) -> list[dict[str, Any]]:
        executable = shutil.which("libreoffice") or shutil.which("soffice")
        if not executable:
            return []
        try:
            with tempfile.TemporaryDirectory(prefix="ledgerlens-docx-") as temp_dir:
                subprocess.run(
                    [executable, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(path)],
                    check=True,
                    capture_output=True,
                    timeout=45,
                )
                rendered = Path(temp_dir) / f"{path.stem}.pdf"
                if not rendered.exists():
                    return []
                locations: list[dict[str, Any]] = []
                with pdfplumber.open(rendered) as pdf:
                    for page_number, page in enumerate(pdf.pages, start=1):
                        words = page.extract_words(use_text_flow=True, keep_blank_chars=False)
                        lines: list[list[dict[str, Any]]] = []
                        for word in words:
                            if not lines or abs(float(word["top"]) - float(lines[-1][0]["top"])) > 3:
                                lines.append([word])
                            else:
                                lines[-1].append(word)
                        for line in lines:
                            line.sort(key=lambda item: float(item["x0"]))
                            locations.append(
                                {
                                    "page": page_number,
                                    "bbox": [
                                        min(float(item["x0"]) for item in line),
                                        min(float(item["top"]) for item in line),
                                        max(float(item["x1"]) for item in line),
                                        max(float(item["bottom"]) for item in line),
                                    ],
                                    "text": " ".join(str(item["text"]) for item in line),
                                }
                            )
                return locations
        except (OSError, subprocess.SubprocessError):
            return []

    def _parse_xml(self, path: Path, bundle: ParseBundle) -> None:
        text = path.read_text(encoding=detect_encoding(path), errors="replace")
        for line_number, line in enumerate(text.splitlines(), start=1):
            visible = re.sub(r"<[^>]+>", " ", line)
            visible = re.sub(r"\s+", " ", visible).strip()
            if visible:
                self._anchor(bundle, AnchorKind.XML, {"line": line_number}, visible)
        role = DatasetRole.GDPDU_SCHEMA if path.name.lower() == "index.xml" else infer_text_role(text)
        bundle.datasets.append(ParsedDataset(path.name, role, [], {}, metadata={"line_count": len(text.splitlines())}))

    def _parse_text(self, path: Path, bundle: ParseBundle, kind: AnchorKind) -> None:
        text = path.read_text(encoding=detect_encoding(path), errors="replace")
        for line_number, line in enumerate(text.splitlines(), start=1):
            if line.strip():
                self._anchor(bundle, kind, {"line": line_number}, line.strip())
        role = infer_text_role(text)
        bundle.datasets.append(ParsedDataset(path.name, role, [], {}, metadata={"line_count": len(text.splitlines())}))
